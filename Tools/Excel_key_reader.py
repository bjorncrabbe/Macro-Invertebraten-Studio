import os

from openpyxl import load_workbook

from excel_node import ExcelNode


class ExcelKeyReader:

    def __init__(self, filename):

        self.filename = filename

        self.workbook = None
        self.sheet = None

        self.nodes = []

    def open(self, sheet_name):

        self.workbook = load_workbook(
            self.filename,
            data_only=True
        )

        self.sheet = self.workbook[sheet_name]

    def print_cells(self):

        for row in self.sheet.iter_rows():

            for cell in row:

                if cell.value is None:
                    continue

                text = str(cell.value).strip()

                if text == "":
                    continue

                print(f"{cell.coordinate:5} | {text}")

    def detect_type(self, text):

        if text.isdigit():
            return "option"

        if text.isupper():
            return "result"

        return "question"

    def read_nodes(self):

        self.nodes.clear()

        for row in self.sheet.iter_rows():

            for cell in row:

                if cell.value is None:
                    continue

                text = str(cell.value).strip()

                if text == "":
                    continue

                node = ExcelNode(
                    coordinate=cell.coordinate,
                    row=cell.row,
                    column=cell.column,
                    text=text,
                    node_type=self.detect_type(text)
                )

                self.nodes.append(node)


if __name__ == "__main__":

    BASE_DIR = os.path.dirname(
        os.path.dirname(__file__)
    )

    EXCEL_FILE = os.path.join(
        BASE_DIR,
        "Data",
        "Determination",
        "excel",
        "DeterminatiesleutelMacroinvertebraten.xlsx"
    )

    reader = ExcelKeyReader(EXCEL_FILE)

    reader.open("HIRUDINEA")

    reader.read_nodes()

    print(f"\nAantal nodes: {len(reader.nodes)}\n")

    for node in reader.nodes:

        print(
            f"{node.coordinate:5} | "
            f"{node.node_type:8} | "
            f"{node.text}"
        )